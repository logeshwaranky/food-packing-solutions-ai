import openpyxl
import pandas as pd
import json

excel_path = r"C:\Users\ELCOT\Downloads\food_packing_ai_dataset.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
print("Sheet names found:", wb.sheetnames)

summary = {}

for name in wb.sheetnames:
    df = pd.read_excel(excel_path, sheet_name=name)
    summary[name] = {
        "shape": df.shape,
        "columns": list(df.columns),
        "null_counts": df.isnull().sum().to_dict(),
        "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
        "head_sample": df.head(3).to_dict(orient="records")
    }

print("\n" + "="*80)
for name, data in summary.items():
    print(f"\nSHEET: {name} | Dimensions: {data['shape']}")
    print("Columns:", data['columns'])
    print("Null Counts:", data['null_counts'])
    print("First 2 rows sample:")
    for row in data['head_sample'][:2]:
        print("  ", row)

with open("dataset_inspection.json", "w", encoding="utf-8") as f:
    json.dump(summary, f, indent=2, default=str)
print("\nFull inspection saved to dataset_inspection.json")
